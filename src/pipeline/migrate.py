"""V1 -> V2 migration utilities.

The V1 workbook stores daily percentage returns. Those values are useful for
regression/audit purposes, but they are NOT sufficient to reconstruct OHLCV.
V2 therefore keeps them in legacy_daily_returns and leaves daily_prices to the
real OHLCV provider.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import json
import math
import shutil
from typing import Iterable

import openpyxl

from src.data.storage.database import Database
from src.data.universe.manager import UniverseManager, UniverseRecord

RAW_SHEET = "BIST D Return Data"
META_SHEET = "MOMENTUM SCREENER"
DATE_ROW = 3
DATA_START_ROW = 6
TICKER_COL = 4  # D


@dataclass(frozen=True)
class MigrationStats:
    workbook_tickers: int = 0
    return_rows: int = 0
    return_cells: int = 0
    json_files: int = 0
    json_return_cells: int = 0


def _finite_float(value: object) -> float | None:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def _date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def import_v1_excel(db: Database, excel_path: str | Path, source_name: str | None = None) -> MigrationStats:
    """Import V1 metadata and daily returns without modifying the workbook."""
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(excel_path)

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = wb[RAW_SHEET]
        meta_ws = wb[META_SHEET] if META_SHEET in wb.sheetnames else None

        metadata: dict[str, dict] = {}
        if meta_ws:
            for row in meta_ws.iter_rows(min_row=5, max_row=meta_ws.max_row, values_only=True):
                ticker = row[4] if len(row) > 4 else None
                if not ticker:
                    continue
                ticker = UniverseManager.normalize_ticker(str(ticker))
                metadata[ticker] = {
                    "type": str(row[1]).strip() if len(row) > 1 and row[1] else "Stock",
                    "industry": str(row[2]).strip() if len(row) > 2 and row[2] else None,
                    "name": str(row[5]).strip() if len(row) > 5 and row[5] else ticker,
                }

        date_row = next(ws.iter_rows(min_row=DATE_ROW, max_row=DATE_ROW, values_only=True))
        dates = {i: _date_value(v) for i, v in enumerate(date_row)}
        dates = {i: d for i, d in dates.items() if d is not None}

        source = source_name or excel_path.name
        imported_at = datetime.now().astimezone().isoformat()
        stats = MigrationStats(workbook_tickers=len(metadata))
        row_count = 0
        cell_count = 0

        for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
            if len(row) <= TICKER_COL - 1 or not row[TICKER_COL - 1]:
                continue
            ticker = UniverseManager.normalize_ticker(str(row[TICKER_COL - 1]))
            meta = metadata.get(ticker, {"name": ticker, "industry": None, "type": "Stock"})
            instrument_type = "EQUITY" if meta["type"].lower() in {"stock", "equity", "hisse", "pay"} else "OTHER"
            security_id = UniverseManager(db.conn).upsert(
                UniverseRecord(ticker=ticker, name=meta["name"], industry=meta["industry"], instrument_type=instrument_type),
                as_of=min((d.isoformat() for d in dates.values()), default=date.today().isoformat()),
            )
            row_count += 1

            for idx, d in dates.items():
                if idx >= len(row):
                    continue
                value = _finite_float(row[idx])
                if value is None:
                    continue
                # V1 stores percentages (e.g. 1.25 = +1.25%).
                ret = value / 100.0
                db.conn.execute(
                    """INSERT OR REPLACE INTO legacy_daily_returns
                       (security_id,ticker,date,return_decimal,source_file,source_type,imported_at)
                       VALUES (?,?,?,?,?,?,?)""",
                    (security_id, ticker, d.isoformat(), ret, source, "V1_EXCEL", imported_at),
                )
                cell_count += 1

        db.conn.commit()
        return MigrationStats(
            workbook_tickers=len(metadata),
            return_rows=row_count,
            return_cells=cell_count,
        )
    finally:
        wb.close()


def import_v1_json(db: Database, json_paths: Iterable[str | Path]) -> MigrationStats:
    """Import V1 detail JSON returns and top-level Momentum/FIP snapshots."""
    paths = [Path(p) for p in json_paths]
    cells = 0
    files = 0
    imported_at = datetime.now().astimezone().isoformat()
    manager = UniverseManager(db.conn)

    for path in paths:
        if not path.exists() or path.suffix.lower() != ".json":
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        ticker = manager.normalize_ticker(str(payload.get("ticker", "")))
        if not ticker:
            continue
        security_id = manager.upsert(
            UniverseRecord(
                ticker=ticker,
                name=payload.get("name") or ticker,
                instrument_type="EQUITY" if str(payload.get("type", "Stock")).lower() == "stock" else "OTHER",
                industry=payload.get("industry"),
            ),
            as_of=min(
                (x.get("date") for m in payload.get("monthly", []) for x in m.get("days", []) if x.get("date")),
                default=date.today().isoformat(),
            ),
        )
        for month in payload.get("monthly", []):
            for day in month.get("days", []):
                if not day.get("date") or day.get("ret") is None:
                    continue
                db.conn.execute(
                    """INSERT OR REPLACE INTO legacy_daily_returns
                       (security_id,ticker,date,return_decimal,source_file,source_type,imported_at)
                       VALUES (?,?,?,?,?,?,?)""",
                    (security_id, ticker, day["date"], float(day["ret"]), path.name, "V1_JSON", imported_at),
                )
                cells += 1
        if payload.get("mom") is not None or payload.get("fip") is not None:
            db.conn.execute(
                """INSERT INTO legacy_indicator_snapshots
                   (ticker,as_of_date,momentum,fip,source_file,imported_at)
                   VALUES (?,?,?,?,?,?)""",
                (ticker, None, payload.get("mom"), payload.get("fip"), path.name, imported_at),
            )
        files += 1

    db.conn.commit()
    return MigrationStats(json_files=files, json_return_cells=cells)


def archive_legacy_file(source: str | Path, archive_dir: str | Path) -> Path:
    """Copy a legacy file into V2 archive without altering the source."""
    source = Path(source)
    archive_dir = Path(archive_dir)
    archive_dir.mkdir(parents=True, exist_ok=True)
    target = archive_dir / source.name
    shutil.copy2(source, target)
    return target
