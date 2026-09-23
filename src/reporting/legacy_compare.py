from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
import calendar as _calendar
import json
import math
from pathlib import Path
from statistics import mean
from typing import Iterable

from openpyxl import load_workbook

from src.indicators.momentum import cumulative_return
from src.indicators.fip import sign


@dataclass(frozen=True)
class LegacyCompareRow:
    ticker: str
    v1_as_of: str
    v1_json_momentum: float | None
    v1_json_fip: float | None
    v1_excel_momentum: float | None
    v1_excel_fip: float | None
    v1_excel_json_match: bool | None
    v1_json_reconstructed_momentum: float | None
    v1_json_reconstructed_fip: float | None
    v1_json_rounding_match: bool | None
    v2_as_of: str | None
    v2_momentum_252: float | None
    v2_fip_252: float | None
    momentum_difference: float | None
    fip_difference: float | None
    methodology_comparable: bool = False


def _safe_diff(a, b):
    if a is None or b is None:
        return None
    return float(b) - float(a)


def _previous_month(year: int, month: int) -> tuple[int, int]:
    return (year - 1, 12) if month == 1 else (year, month - 1)


def _shift_month(year: int, month: int, delta: int) -> tuple[int, int]:
    absolute = year * 12 + (month - 1) + delta
    return absolute // 12, absolute % 12 + 1


def _legacy_window(as_of: date) -> tuple[date, date]:
    end_year, end_month = _previous_month(as_of.year, as_of.month)
    start_year, start_month = _shift_month(end_year, end_month, -11)
    start = date(start_year, start_month, 1)
    end = date(end_year, end_month, _calendar.monthrange(end_year, end_month)[1])
    return start, end


def _infer_json_as_of(paths: Iterable[Path]) -> date:
    latest: date | None = None
    for path in paths:
        payload = json.loads(path.read_text(encoding="utf-8"))
        for month in payload.get("monthly") or []:
            for day in month.get("days") or []:
                raw = day.get("date")
                if not raw:
                    continue
                d = date.fromisoformat(raw)
                if latest is None or d > latest:
                    latest = d
    if latest is None:
        raise ValueError("Could not infer V1 as-of date from detail JSONs")
    return latest


def _json_daily(payload: dict, start: date, end: date) -> dict[date, float]:
    out: dict[date, float] = {}
    for month in payload.get("monthly") or []:
        for day in month.get("days") or []:
            raw = day.get("date")
            if raw is None or day.get("ret") is None:
                continue
            d = date.fromisoformat(raw)
            if start <= d <= end:
                out[d] = float(day["ret"])
    return out


def _date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    if isinstance(value, (int, float)) and 40000 < value < 60000:
        return date(1899, 12, 30) + __import__('datetime').timedelta(days=int(value))
    return None


def _finite_float(value: object) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def load_v1_excel_returns(excel_path: str | Path) -> tuple[dict[str, dict[date, float]], list[date]]:
    """Read V1 raw daily-return sheet without mutating the workbook."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = workbook["BIST D Return Data"]
        date_row = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
        idx_to_date = {i: _date_value(value) for i, value in enumerate(date_row)}
        idx_to_date = {i: d for i, d in idx_to_date.items() if d is not None}
        daily: dict[str, dict[date, float]] = {}
        for row in ws.iter_rows(min_row=6, max_row=650, values_only=True):
            if len(row) < 4 or not row[3]:
                continue
            ticker = str(row[3]).strip().upper().replace(".IS", "")
            values: dict[date, float] = {}
            for index, d in idx_to_date.items():
                if index >= len(row):
                    continue
                number = _finite_float(row[index])
                if number is not None:
                    values[d] = number / 100.0
            daily[ticker] = values
        return daily, sorted(idx_to_date.values())
    finally:
        workbook.close()


def _calc_legacy(daily: dict[date, float], start: date, end: date, denominator: int) -> tuple[float | None, float | None]:
    values = [r for d, r in daily.items() if start <= d <= end]
    if len(values) < 10:
        return None, None
    momentum = cumulative_return(values)
    if momentum is None:
        return None, None
    neg = sum(r < 0 for r in values)
    pos = sum(r > 0 for r in values)
    fip = sign(momentum) * (neg - pos) / denominator
    return round(momentum, 6), round(fip, 6)


def _json_rounding_tolerance(momentum: float | None) -> float:
    # Daily returns in V1 detail JSON are rounded to 6 decimals whereas top-level
    # momentum was calculated from full-precision Excel returns. The error grows
    # with both compounding and momentum magnitude; this bound is audit-only.
    if momentum is None:
        return 1e-6
    return max(1e-4, abs(momentum) * 1e-5)


def build_legacy_comparison(
    *,
    v1_detail_dir: str | Path,
    v1_excel_path: str | Path | None = None,
    v2_screener_json: str | Path | None = None,
    as_of: date | None = None,
) -> dict:
    """Build the three-way V1 Excel ↔ V1 JSON ↔ V2 audit report.

    PASS/FAIL uses the authoritative V1 Excel-to-V1-JSON comparison when the
    workbook is provided. V1-vs-V2 equality is never required because V1 uses
    a legacy 12-1 calendar-month window while V2 uses fixed trading-day windows.
    """
    detail_dir = Path(v1_detail_dir)
    paths = sorted(detail_dir.glob("*.json"))
    if not paths:
        raise ValueError("No V1 detail JSON files found")
    if as_of is None:
        as_of = _infer_json_as_of(paths)
    start, end = _legacy_window(as_of)

    payloads = {path.stem.upper(): json.loads(path.read_text(encoding="utf-8")) for path in paths}
    json_union_days = {
        d for payload in payloads.values() for d in _json_daily(payload, start, end).keys()
    }
    json_denominator = len(json_union_days)
    if json_denominator == 0:
        raise ValueError("No V1 JSON trading days found inside the 12-1 window")

    excel_daily: dict[str, dict[date, float]] = {}
    excel_header_dates: list[date] = []
    excel_union_days: set[date] = set()
    if v1_excel_path:
        excel_daily, excel_header_dates = load_v1_excel_returns(v1_excel_path)
        excel_union_days = {
            d for series in excel_daily.values() for d in series.keys() if start <= d <= end
        }
        if not excel_union_days:
            raise ValueError("No V1 Excel trading days found inside the 12-1 window")

    v2_by_ticker: dict[str, dict] = {}
    if v2_screener_json:
        v2_rows = json.loads(Path(v2_screener_json).read_text(encoding="utf-8"))
        v2_by_ticker = {str(row.get("ticker", "")).upper(): row for row in v2_rows}

    rows_out: list[LegacyCompareRow] = []
    failures: list[str] = []
    json_rounding_mismatches: list[str] = []

    for ticker, payload in sorted(payloads.items()):
        expected_mom = payload.get("mom")
        expected_fip = payload.get("fip")

        json_series = _json_daily(payload, start, end)
        json_mom, json_fip = _calc_legacy(json_series, start, end, json_denominator)
        json_rounding_match = (
            (expected_mom is None and json_mom is None)
            or (
                expected_mom is not None
                and json_mom is not None
                and abs(float(expected_mom) - float(json_mom)) <= _json_rounding_tolerance(float(expected_mom))
            )
        ) and (
            (expected_fip is None and json_fip is None)
            or (expected_fip is not None and json_fip is not None and abs(float(expected_fip) - float(json_fip)) <= 1e-6)
        )
        if not json_rounding_match:
            json_rounding_mismatches.append(ticker)

        excel_mom = excel_fip = None
        excel_match: bool | None = None
        if v1_excel_path:
            series = excel_daily.get(ticker)
            if series is None:
                excel_match = False
                failures.append(f"{ticker}: missing from V1 Excel")
            else:
                excel_mom, excel_fip = _calc_legacy(series, start, end, len(excel_union_days))
                excel_match = excel_mom == expected_mom and excel_fip == expected_fip
                if not excel_match:
                    failures.append(
                        f"{ticker}: Excel/JSON mismatch "
                        f"mom {excel_mom!r}!={expected_mom!r}, fip {excel_fip!r}!={expected_fip!r}"
                    )

        v2 = v2_by_ticker.get(ticker, {})
        rows_out.append(
            LegacyCompareRow(
                ticker=ticker,
                v1_as_of=as_of.isoformat(),
                v1_json_momentum=expected_mom,
                v1_json_fip=expected_fip,
                v1_excel_momentum=excel_mom,
                v1_excel_fip=excel_fip,
                v1_excel_json_match=excel_match,
                v1_json_reconstructed_momentum=json_mom,
                v1_json_reconstructed_fip=json_fip,
                v1_json_rounding_match=json_rounding_match,
                v2_as_of=v2.get("as_of"),
                v2_momentum_252=v2.get("momentum_252"),
                v2_fip_252=v2.get("fip_252"),
                momentum_difference=_safe_diff(expected_mom, v2.get("momentum_252")),
                fip_difference=_safe_diff(expected_fip, v2.get("fip_252")),
            )
        )

    comparable = [r for r in rows_out if r.v2_as_of]
    abs_mom = [abs(r.momentum_difference) for r in comparable if r.momentum_difference is not None]
    abs_fip = [abs(r.fip_difference) for r in comparable if r.fip_difference is not None]
    excel_matches = [r for r in rows_out if r.v1_excel_json_match is True]

    # If Excel is present, it is the authoritative regression gate. Otherwise
    # the rounded JSON reconstruction is the best available fallback.
    if v1_excel_path:
        status = "PASS" if not failures and len(excel_matches) == len(rows_out) else "FAIL"
    else:
        status = "PASS" if not json_rounding_mismatches else "WARN"

    future_excel_dates = [d for d in excel_header_dates if d > as_of]
    return {
        "schema_version": 2,
        "status": status,
        "v1_as_of": as_of.isoformat(),
        "legacy_window": {"start": start.isoformat(), "end": end.isoformat()},
        "note": (
            "V1 Excel is the legacy source of truth. V1 detail JSON rounds daily returns to 6 decimals, "
            "so JSON self-reconstruction is informational. V1 uses a legacy 12-1 calendar-month window "
            "while V2 uses fixed trading-day windows; V1-vs-V2 differences are descriptive only."
        ),
        "summary": {
            "v1_files_checked": len(paths),
            "v1_excel_json_matches": len(excel_matches) if v1_excel_path else None,
            "v1_json_rounding_matches": sum(r.v1_json_rounding_match is True for r in rows_out),
            "v1_json_rounding_mismatches": len(json_rounding_mismatches),
            "v1_json_market_days": json_denominator,
            "v1_excel_market_days": len(excel_union_days) if v1_excel_path else None,
            "v1_excel_future_header_dates_after_as_of": len(future_excel_dates) if v1_excel_path else None,
            "v1_excel_last_header_date": max(excel_header_dates).isoformat() if excel_header_dates else None,
            "v2_common_tickers": len(comparable),
            "mean_absolute_momentum_difference": mean(abs_mom) if abs_mom else None,
            "mean_absolute_fip_difference": mean(abs_fip) if abs_fip else None,
            "failures": len(failures),
        },
        "failures": failures,
        "json_rounding_mismatch_tickers": json_rounding_mismatches,
        "rows": [asdict(r) for r in rows_out],
    }


def write_legacy_comparison(report: dict, output: str | Path) -> Path:
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
