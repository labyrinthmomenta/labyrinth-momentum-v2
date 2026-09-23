from __future__ import annotations

"""Official-source BIST equity universe acquisition.

Current equity identities come from KAP's market page (linked by Borsa Istanbul's
"İşlem Gören Şirketler" page). First-trade dates are enriched from Borsa
Istanbul's documented ``ilkislem.zip`` report. The hidden report URL is
discovered from Borsa Istanbul's official ``DataFilePaths.zip`` catalogue so
we do not depend on an undocumented static path.

All network/file-format failures are fail-closed: a broken upstream source must
not silently shrink or corrupt the canonical universe.
"""

from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from io import BytesIO
import csv
import io
import re
from pathlib import PurePosixPath
import urllib.parse
import urllib.request
import zipfile
from typing import Callable, Iterable

from openpyxl import load_workbook

from .manager import UniverseRecord, UniverseManager

KAP_MARKETS_URL = "https://www.kap.org.tr/tr/Pazarlar"
BIST_DATA_PATHS_URL = "https://www.borsaistanbul.com/files/DataFilePaths.zip"

# Company-share markets. Instrument/fund-only markets are intentionally omitted.
EQUITY_MARKETS = {
    "YILDIZ PAZAR",
    "ANA PAZAR",
    "ALT PAZAR",
    "YAKIN İZLEME PAZARI",
    "PİYASA ÖNCESİ İŞLEM PLATFORMU",
}

_TICKER_RE = re.compile(r"^[A-Z0-9]{2,12}$")
_URL_RE = re.compile(r"https?://[^\s\"'<>]+", re.IGNORECASE)


class OfficialUniverseError(RuntimeError):
    pass


@dataclass(frozen=True)
class OfficialUniverseSnapshot:
    records: list[UniverseRecord]
    kap_records: int
    first_trade_records: int
    first_trade_url: str
    source_date: str
    market_counts: dict[str, int]


@dataclass(frozen=True)
class MarketRow:
    ticker: str
    name: str
    market: str


class _TableParser(HTMLParser):
    """Small dependency-free HTML table parser for server-rendered KAP rows."""

    def __init__(self):
        super().__init__()
        self.in_tr = False
        self.in_cell = False
        self.cell_parts: list[str] = []
        self.row: list[str] = []
        self.rows: list[list[str]] = []

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "tr":
            self.in_tr = True
            self.row = []
        elif self.in_tr and tag in {"td", "th"}:
            self.in_cell = True
            self.cell_parts = []

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self.in_tr and tag in {"td", "th"} and self.in_cell:
            text = " ".join("".join(self.cell_parts).split())
            self.row.append(text)
            self.in_cell = False
        elif tag == "tr" and self.in_tr:
            if self.row:
                self.rows.append(self.row)
            self.in_tr = False
            self.row = []

    def handle_data(self, data):
        if self.in_cell:
            self.cell_parts.append(data)


def _clean_market(text: str) -> str | None:
    upper = " ".join(text.upper().split())
    for market in EQUITY_MARKETS:
        if market in upper:
            return market
    # Explicit non-equity headings reset the parser so rows beneath them cannot
    # inherit the previous equity market.
    if "PAZAR" in upper or "PİYASASI" in upper or "PLATFORM" in upper:
        return "__OTHER__"
    return None


def parse_kap_markets_html(html: str) -> list[MarketRow]:
    parser = _TableParser()
    parser.feed(html)
    current_market: str | None = None
    output: list[MarketRow] = []
    seen: set[str] = set()

    for cells in parser.rows:
        joined = " ".join(cells)
        market = _clean_market(joined)
        if market is not None and len(cells) <= 3:
            current_market = market
            # A heading row can still carry a company row in unusual HTML; do
            # not continue until ticker extraction below has had a chance.

        # Typical KAP row: [Sıra, Kod, Şirket / Fon Adı]. Tolerate extra cells.
        ticker = None
        ticker_index = None
        for index, raw in enumerate(cells):
            candidate = raw.strip().upper().replace(".IS", "")
            if _TICKER_RE.fullmatch(candidate) and not candidate.isdigit():
                ticker = candidate
                ticker_index = index
                break
        if not ticker or ticker in {"KOD", "SIRA"}:
            continue
        if current_market not in EQUITY_MARKETS:
            continue

        name = ""
        for raw in cells[(ticker_index or 0) + 1 :]:
            value = " ".join(raw.split())
            if value and value.upper() not in {"ŞİRKET / FON ADI", "ŞİRKET", "FON"}:
                name = value
                break
        if not name:
            continue
        if ticker in seen:
            raise OfficialUniverseError(f"Duplicate ticker in KAP market source: {ticker}")
        seen.add(ticker)
        output.append(MarketRow(ticker=ticker, name=name, market=current_market))

    return output


def _iter_text_from_xlsx(data: bytes) -> Iterable[str]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        yield str(value)
    finally:
        workbook.close()


def _iter_catalog_strings(data: bytes) -> Iterable[str]:
    """Yield strings from DataFilePaths.zip without assuming its inner format."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.endswith("/"):
                continue
            raw = archive.read(name)
            suffix = PurePosixPath(name).suffix.lower()
            if suffix in {".xlsx", ".xlsm"}:
                yield from _iter_text_from_xlsx(raw)
            else:
                for encoding in ("utf-8-sig", "cp1254", "latin-1"):
                    try:
                        text = raw.decode(encoding)
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    continue
                yield from text.splitlines()


def discover_ilkislem_url(catalog_zip: bytes, *, base_url: str = BIST_DATA_PATHS_URL) -> str:
    candidates: list[str] = []
    for text in _iter_catalog_strings(catalog_zip):
        if "ilkislem" not in text.lower():
            continue
        candidates.extend(_URL_RE.findall(text))
        # Semicolon/tab separated catalogues sometimes provide a relative path.
        for token in re.split(r"[;\t, ]+", text):
            cleaned = token.strip().strip('"\'')
            if "ilkislem" in cleaned.lower() and cleaned.lower().endswith(".zip"):
                candidates.append(cleaned)
    for candidate in candidates:
        url = urllib.parse.urljoin(base_url, candidate)
        if "ilkislem" in url.lower() and url.lower().endswith(".zip"):
            return url
    raise OfficialUniverseError("BIST DataFilePaths catalogue does not expose ilkislem.zip")


def _parse_date(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def parse_first_trade_zip(data: bytes) -> dict[str, str]:
    """Parse BIST ilkislem.zip -> current ticker -> first trading date."""
    result: dict[str, str] = {}
    with zipfile.ZipFile(BytesIO(data)) as archive:
        members = [n for n in archive.namelist() if not n.endswith("/")]
        if not members:
            raise OfficialUniverseError("ilkislem.zip is empty")
        target = next((n for n in members if n.lower().endswith((".xlsx", ".xlsm"))), None)
        if target is not None:
            workbook = load_workbook(BytesIO(archive.read(target)), read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        values = list(row)
                        if len(values) < 5:
                            continue
                        ticker = str(values[1] or "").strip().upper().replace(".IS", "")
                        first_day = _parse_date(values[4])
                        if _TICKER_RE.fullmatch(ticker) and first_day:
                            result[ticker] = first_day
            finally:
                workbook.close()
        else:
            target = members[0]
            raw = archive.read(target)
            text = None
            for encoding in ("utf-8-sig", "cp1254", "latin-1"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    pass
            if text is None:
                raise OfficialUniverseError("Unsupported ilkislem.zip inner format")
            reader = csv.reader(io.StringIO(text), delimiter=";")
            for row in reader:
                if len(row) < 5:
                    continue
                ticker = row[1].strip().upper().replace(".IS", "")
                first_day = _parse_date(row[4])
                if _TICKER_RE.fullmatch(ticker) and first_day:
                    result[ticker] = first_day
    if not result:
        raise OfficialUniverseError("No first-trade records parsed from ilkislem.zip")
    return result


def default_http_get(url: str, timeout: int = 30) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "LabyrinthMomentumV2/1.0 (+https://github.com/)"},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read()


def fetch_official_universe(
    *,
    http_get: Callable[[str], bytes] = default_http_get,
    kap_url: str = KAP_MARKETS_URL,
    data_paths_url: str = BIST_DATA_PATHS_URL,
    min_equities: int = 400,
) -> OfficialUniverseSnapshot:
    try:
        kap_html = http_get(kap_url).decode("utf-8", errors="replace")
        market_rows = parse_kap_markets_html(kap_html)
        if len(market_rows) < min_equities:
            raise OfficialUniverseError(
                f"KAP equity universe unexpectedly small: {len(market_rows)} < {min_equities}"
            )
        catalogue = http_get(data_paths_url)
        first_trade_url = discover_ilkislem_url(catalogue, base_url=data_paths_url)
        first_trade = parse_first_trade_zip(http_get(first_trade_url))
    except OfficialUniverseError:
        raise
    except Exception as exc:
        raise OfficialUniverseError(f"Official universe download failed: {exc}") from exc

    market_counts: dict[str, int] = {}
    for row in market_rows:
        market_counts[row.market] = market_counts.get(row.market, 0) + 1

    records = [
        UniverseRecord(
            ticker=row.ticker,
            name=row.name,
            instrument_type="EQUITY",
            first_trade_date=first_trade.get(row.ticker),
            status="ACTIVE",
            active=1,
        )
        for row in market_rows
    ]
    return OfficialUniverseSnapshot(
        records=records,
        kap_records=len(market_rows),
        first_trade_records=len(first_trade),
        first_trade_url=first_trade_url,
        source_date=date.today().isoformat(),
        market_counts=market_counts,
    )


def guard_universe_change(
    conn,
    records: Iterable[UniverseRecord],
    *,
    max_drop_fraction: float = 0.08,
    min_equities: int = 400,
) -> None:
    """Fail closed on suspicious source shrinkage before marking anything inactive."""
    records = list(records)
    incoming = {UniverseManager.normalize_ticker(r.ticker) for r in records if r.instrument_type == "EQUITY" and r.active}
    if len(incoming) < min_equities:
        raise OfficialUniverseError(f"Incoming equity universe below safety floor: {len(incoming)}")
    current = {
        row[0]
        for row in conn.execute(
            """SELECT si.ticker FROM securities s
               JOIN security_identifiers si ON si.security_id=s.security_id
               WHERE s.instrument_type='EQUITY' AND s.active=1 AND si.is_current=1"""
        ).fetchall()
    }
    if current:
        missing = current - incoming
        fraction = len(missing) / len(current)
        if fraction > max_drop_fraction:
            raise OfficialUniverseError(
                f"Suspicious universe shrinkage: {len(missing)}/{len(current)} active equities missing "
                f"({fraction:.1%} > {max_drop_fraction:.1%})"
            )
