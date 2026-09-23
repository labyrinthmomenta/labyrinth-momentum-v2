from __future__ import annotations

from datetime import date
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook

from src.reporting.legacy_compare import build_legacy_comparison, write_legacy_comparison


def _make_v1_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'BIST D Return Data'
    # Sep 2025 -> Aug 2026 window plus future header columns that must not move as-of.
    days = [date(2025, 9, 1), date(2025, 9, 2), date(2025, 9, 3), date(2025, 9, 4), date(2025, 9, 5),
            date(2025, 9, 8), date(2025, 9, 9), date(2025, 9, 10), date(2025, 9, 11), date(2025, 9, 12)]
    for idx, d in enumerate(days, 5):
        ws.cell(3, idx).value = d
    ws.cell(3, 15).value = date(2026, 12, 31)
    ws.cell(6, 4).value = 'AAA'
    returns_pct = [1.0, -1.0, 1.0, -1.0, 1.0, -1.0, 1.0, -1.0, 1.0, -1.0]
    for idx, value in enumerate(returns_pct, 5):
        ws.cell(6, idx).value = value
    ws.cell(6, 15).value = 0.0
    ms = wb.create_sheet('MOMENTUM SCREENER')
    ms.cell(5, 5).value = 'AAA'
    ms.cell(5, 6).value = 'Alpha'
    wb.save(path)
    wb.close()


def _make_detail(path: Path):
    rets = [0.01, -0.01, 0.01, -0.01, 0.01, -0.01, 0.01, -0.01, 0.01, -0.01]
    product = 1.0
    for r in rets:
        product *= 1 + r
    payload = {
        'ticker': 'AAA', 'mom': round(product - 1, 6), 'fip': 0.0,
        'monthly': [
            {'month': '2025-09', 'days': [
                {'date': d.isoformat(), 'ret': r}
                for d, r in zip([
                    date(2025, 9, 1), date(2025, 9, 2), date(2025, 9, 3), date(2025, 9, 4), date(2025, 9, 5),
                    date(2025, 9, 8), date(2025, 9, 9), date(2025, 9, 10), date(2025, 9, 11), date(2025, 9, 12)
                ], rets)
            ]},
            {'month': '2026-09', 'days': [{'date': '2026-09-21', 'ret': 0.0}]},
        ]
    }
    path.write_text(json.dumps(payload), encoding='utf-8')


def test_legacy_report_uses_json_as_of_and_exact_excel_regression(tmp_path):
    detail = tmp_path / 'detail'; detail.mkdir()
    _make_detail(detail / 'AAA.json')
    xlsx = tmp_path / 'v1.xlsx'; _make_v1_workbook(xlsx)
    screener = tmp_path / 'screener.json'
    screener.write_text(json.dumps([{'ticker': 'AAA', 'as_of': '2026-09-21', 'momentum_252': 0.1, 'fip_252': 0.2}]), encoding='utf-8')
    report = build_legacy_comparison(v1_detail_dir=detail, v1_excel_path=xlsx, v2_screener_json=screener)
    assert report['status'] == 'PASS'
    assert report['v1_as_of'] == '2026-09-21'
    assert report['summary']['v1_excel_json_matches'] == 1
    assert report['summary']['v1_excel_future_header_dates_after_as_of'] == 1
    assert report['rows'][0]['methodology_comparable'] is False
    out = write_legacy_comparison(report, tmp_path / 'report.json')
    assert out.exists()
