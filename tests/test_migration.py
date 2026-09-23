from datetime import date
import json
from pathlib import Path
import sqlite3
import openpyxl

from src.data.storage.database import Database
from src.pipeline.migrate import import_v1_excel, import_v1_json


def make_db(tmp_path):
    db = Database(tmp_path / "test.db")
    db.initialize()
    return db


def test_excel_returns_are_imported_as_legacy_not_prices(tmp_path):
    xlsx = tmp_path / "v1.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BIST D Return Data"
    ws.cell(3, 5).value = date(2026, 1, 2)
    ws.cell(3, 6).value = date(2026, 1, 5)
    ws.cell(6, 4).value = "TEST"
    ws.cell(6, 5).value = 1.25
    ws.cell(6, 6).value = -2.5
    ms = wb.create_sheet("MOMENTUM SCREENER")
    ms.cell(5, 2).value = "Stock"
    ms.cell(5, 5).value = "TEST"
    ms.cell(5, 6).value = "TEST COMPANY"
    wb.save(xlsx)

    db = make_db(tmp_path)
    stats = import_v1_excel(db, xlsx)
    assert stats.return_cells == 2
    rows = db.conn.execute("SELECT ticker,date,return_decimal FROM legacy_daily_returns ORDER BY date").fetchall()
    assert [(r[0], r[1], r[2]) for r in rows] == [
        ("TEST", "2026-01-02", 0.0125),
        ("TEST", "2026-01-05", -0.025),
    ]
    assert db.conn.execute("SELECT COUNT(*) FROM daily_prices").fetchone()[0] == 0


def test_a1cap_json_preserves_known_v1_snapshot(tmp_path):
    fixture = Path(__file__).parent / "fixtures" / "A1CAP.json"
    db = make_db(tmp_path)
    stats = import_v1_json(db, [fixture])
    assert stats.json_files == 1
    row = db.conn.execute(
        "SELECT ticker,momentum,fip FROM legacy_indicator_snapshots WHERE ticker='A1CAP' ORDER BY id DESC LIMIT 1"
    ).fetchone()
    assert row[0] == "A1CAP"
    assert abs(row[1] - (-0.156294)) < 1e-6
    assert abs(row[2] - (-0.099602)) < 1e-6
